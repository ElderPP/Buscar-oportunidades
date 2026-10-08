#!/usr/bin/env python3
"""Gera a planilha acumulada do radar a partir de oportunidades.csv.
Uso: python3 planilha.py  ->  relatorios/Radar de oportunidades AAAA-MM-DD.xlsx"""
import csv
import datetime as dt
import pathlib

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = pathlib.Path(__file__).parent
HOJE = dt.date.today()
DIAS = ["seg", "ter", "qua", "qui", "sex", "sáb", "dom"]
STATUS = ["radar", "interesse", "inscrito", "proposta enviada", "em andamento",
          "aprovado", "reprovado", "perdi prazo", "descartei"]

# (cabeçalho, campo do CSV, largura)
COLUNAS = [
    ("Encontrada em", "encontrada_em", 12),
    ("Trilha", "trilha", 13),
    ("Nota", "nota", 6),
    ("Oportunidade", "titulo", 40),
    ("Órgão / empresa", "orgao", 22),
    ("Local", "local", 18),
    ("Prazo", "prazo", 17),
    ("O que vence", "o_que_vence", 24),
    ("Dias restantes", None, 9),
    ("Valor / salário", "valor", 26),
    ("Cota PCD", "cota_pcd", 24),
    ("O que pode eliminar", "o_que_pode_eliminar", 38),
    ("Documento a providenciar", "documento", 32),
    ("Próximo passo", "proximo_passo", 32),
    ("Link oficial", "link", 30),
    ("Verificado na fonte?", "verificado", 11),
    ("Status (você preenche)", "status", 16),
    ("Minhas anotações", "notas", 34),
]
TRILHAS = {"concursos": "Concurso", "sistema_s": "Sistema S", "contratacao_publica": "Contratação pública",
           "bolsas": "Bolsa", "emprego": "Emprego privado", "pj_privado": "PJ privado"}

FONTE = "Arial"
CAB_FILL = PatternFill("solid", start_color="1F2937")
EDITAVEL = PatternFill("solid", start_color="FFF2CC")
BORDA = Border(bottom=Side(style="thin", color="D9D9D9"))


def data(txt):
    try:
        return dt.date.fromisoformat(txt)
    except (TypeError, ValueError):
        return None


EM_ANDAMENTO = {"inscrito", "proposta enviada", "em andamento"}
ENCERRADOS = {"descartei", "perdi prazo", "reprovado"}


def mostrar(r):
    # Fora da planilha: o que ele descartou, perdeu ou já foi encerrado, e o que venceu sem ele ter entrado (fica só no funil).
    if r["status"] in ENCERRADOS:
        return False
    p = data(r["prazo"])
    return p is None or p >= HOJE or r["status"] in EM_ANDAMENTO


linhas = [r for r in csv.DictReader(open(BASE / "oportunidades.csv", encoding="utf-8")) if mostrar(r)]
# Prazos abertos primeiro (mais urgente no topo), depois sem prazo, depois processos em andamento já sem prazo de inscrição.
def ordem(r):
    p = data(r["prazo"])
    if p is None:
        return (1, dt.date.max)
    return (0, p) if p >= HOJE else (2, -p.toordinal())
linhas.sort(key=ordem)

wb = Workbook()
ws = wb.active
ws.title = "Oportunidades"

for c, (titulo, _, larg) in enumerate(COLUNAS, 1):
    cel = ws.cell(row=1, column=c, value=titulo)
    cel.font = Font(name=FONTE, bold=True, color="FFFFFF")
    cel.fill = CAB_FILL
    cel.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[get_column_letter(c)].width = larg
ws.row_dimensions[1].height = 32

for r, linha in enumerate(linhas, 2):
    for c, (titulo, campo, _) in enumerate(COLUNAS, 1):
        if campo is None:  # Dias restantes, a partir da data do prazo (não do texto da coluna G)
            p = data(linha["prazo"])
            valor = f"=DATE({p.year},{p.month},{p.day})-TODAY()" if p else ""
        elif campo in ("encontrada_em", "prazo"):
            # Texto, para qualquer app mostrar a data igual (alguns trocam dia/mês ou exibem número).
            d = data(linha[campo])
            valor = f"{d:%d/%m/%Y} ({DIAS[d.weekday()]})" if d and campo == "prazo" else (f"{d:%d/%m/%Y}" if d else "sem data")
        elif campo == "nota":
            valor = int(linha[campo]) if linha[campo] else "–"
        elif campo == "trilha":
            valor = TRILHAS.get(linha[campo], linha[campo])
        else:
            valor = linha[campo] or None
        cel = ws.cell(row=r, column=c, value=valor)
        cel.font = Font(name=FONTE, size=10)
        cel.alignment = Alignment(wrap_text=True, vertical="top")
        cel.border = BORDA
        if campo is None:
            cel.number_format = '0;"vencido";"hoje"'
            cel.alignment = Alignment(horizontal="center", vertical="top")
        if campo == "nota":
            cel.alignment = Alignment(horizontal="center", vertical="top")
        if campo == "link" and linha["link"]:
            cel.hyperlink = linha["link"]
            cel.value = "abrir edital"
            cel.font = Font(name=FONTE, size=10, color="0B57D0", underline="single")
        if campo in ("status", "notas"):
            cel.fill = EDITAVEL

ultima = max(len(linhas) + 1, 2)
faixa_dias = f"I2:I{ultima}"
ws.conditional_formatting.add(faixa_dias, FormulaRule(
    formula=['AND(ISNUMBER(I2),I2<0)'], font=Font(color="9CA3AF")))
ws.conditional_formatting.add(faixa_dias, FormulaRule(
    formula=['AND(ISNUMBER(I2),I2>=0,I2<=3)'], fill=PatternFill("solid", start_color="F8CBAD"),
    font=Font(bold=True, color="9C0006")))
ws.conditional_formatting.add(faixa_dias, FormulaRule(
    formula=['AND(ISNUMBER(I2),I2>3,I2<=7)'], fill=PatternFill("solid", start_color="FFE699")))
for r, linha in enumerate(linhas, 2):  # novidades do dia em negrito
    if data(linha["encontrada_em"]) == HOJE:
        for c in range(1, 9):
            ws.cell(row=r, column=c).font = Font(name=FONTE, size=10, bold=True)

dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"Q2:Q{ultima + 200}")

ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{ultima}"

# Aba de instruções
ajuda = wb.create_sheet("Como usar")
textos = [
    ("Radar de oportunidades — como usar", True),
    (f"Gerada em {HOJE.strftime('%d/%m/%Y')}. Cada envio traz tudo o que ainda está aberto, com as novidades do dia em negrito.", False),
    ("", False),
    ("Colunas em amarelo são suas: Status (escolha na lista) e Minhas anotações.", False),
    ("Dias restantes é calculado sozinho a partir do Prazo: vermelho = até 3 dias, amarelo = até 7 dias, cinza = vencido.", False),
    ("A lista vem ordenada: prazos abertos primeiro (mais urgente no topo), depois os sem prazo confirmado.", False),
    ("Saem da lista: o que você descartou, perdeu o prazo ou foi reprovado, e o que venceu sem você ter se inscrito. Processos em que você está inscrito continuam.", False),
    ("Use o filtro do cabeçalho para ver só uma trilha, só nota 4 e 5, ou só o que está com status 'inscrito'.", False),
    ("", False),
    ("Importante: cada planilha nova é gerada do zero. Para seu status e anotações aparecerem nas próximas,", False),
    ("me avise no chat (ex.: 'me inscrevi na UFERSA') que eu registro no funil.", False),
    ("", False),
    ("Verificado na fonte? = 'sim' quando prazo e requisitos foram conferidos no edital oficial; 'não' quando o site bloqueou e é preciso conferir.", False),
    ("Nota de 0 a 5: 5 = atende tudo e a concorrência é baixa; 3 = atende os obrigatórios, sem diferencial; '–' = não deu para avaliar.", False),
]
for i, (t, negrito) in enumerate(textos, 1):
    ajuda.cell(row=i, column=1, value=t).font = Font(name=FONTE, size=12 if negrito else 10, bold=negrito)
ajuda.column_dimensions["A"].width = 120

wb.calculation.fullCalcOnLoad = True  # "Dias restantes" usa HOJE(); recalcula ao abrir
saida = BASE / "relatorios" / f"Radar de oportunidades {HOJE.isoformat()}.xlsx"
wb.save(saida)
print(saida)
